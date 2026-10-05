"""Rebuild presentation artifacts from measured workbook data; no model evaluation."""
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def records(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def finalize_artifacts(work_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    root = Path(work_dir)
    workbook = openpyxl.load_workbook(root / "results.xlsx")
    backbones = records(workbook["Backbones"])
    training = records(workbook["Training"])
    inference = records(workbook["Inference"])
    backbone_lookup = {r["backbone"]: r for r in backbones}
    baseline = next(r for r in training if r["exp_id"] == "T00")
    selected = next(r for r in training if r["exp_id"] == "T08")
    combined = []
    for stage, rows in [("Backbone", backbones), ("Training", training)]:
        for r in rows:
            # B03/T00 are the same baseline; T08/I00 have the same 1-view pipeline.
            if r["exp_id"] in {"B03", "T08"}:
                continue
            reference = backbone_lookup[r["backbone"]]
            combined.append([
                r["exp_id"], stage, r["backbone"],
                r.get("description", "Basic + CE"), "1-view FP32", "0",
                r["macro_f1_val"], r["top1_val"],
                100 * (r["macro_f1_val"] - baseline["macro_f1_val"]),
                r["params_M"], r["gmac"], reference["latency_p95_ms"],
                "Measured B checkpoint" if stage == "Backbone" else
                "Reference B checkpoint; recipe latency not remeasured",
            ])
    for r in inference:
        combined.append([
            r["exp_id"], "Inference", selected["backbone"],
            "T08: CutMix + label smoothing 0.1", r["method"], "0",
            r["macro_f1_val"], r["top1_val"],
            100 * (r["macro_f1_val"] - baseline["macro_f1_val"]),
            selected["params_M"], selected["gmac"], r["p95_ms"],
            "Measured T08 checkpoint; GMAC is one 224 view, excludes TTA cost",
        ])
    combined.sort(key=lambda r: (-r[6], r[0]))
    sheet = workbook["Summary"]
    sheet.delete_rows(1, sheet.max_row)
    sheet.append([
        "exp_id", "stage", "backbone", "recipe", "inference", "seed",
        "macro_f1_val (0-1)", "top1_val (0-1)",
        "delta_vs_T00 (percentage points)", "params (M)",
        "GMAC single 224 view", "p95 batch1 (ms)", "latency / GMAC note",
    ])
    for row in combined[:10]:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 42
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(size=10)
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
        if row[0].value in {"I03", "T00"}:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="E2F0D9")
    for i, width in enumerate([12, 14, 25, 35, 27, 8, 20, 20, 23, 16, 20, 20, 55], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = sheet.dimensions
    workbook.save(root / "results.xlsx")

    figure, axes = plt.subplots(1, 2, figsize=(12, 4.5), layout="constrained")
    for axis, field, label in [
        (axes[0], "latency_p95_ms", "Measured p95 latency, batch 1 (ms)"),
        (axes[1], "params_M", "Parameters (million)"),
    ]:
        for r in backbones:
            axis.scatter(r[field], r["macro_f1_val"], color="#305496", s=45)
            axis.annotate(r["exp_id"], (r[field], r["macro_f1_val"]),
                          xytext=(5, 5), textcoords="offset points")
        axis.set_xlabel(label)
        axis.set_ylabel("Validation macro-F1 (0-1)")
        axis.set_title("Backbone accuracy / cost comparison")
        axis.grid(alpha=0.2)
        axis.margins(0.18)
    figure.suptitle("DeepWeeds fold 0, seed 0, common 12-epoch recipe")
    figure.text(0.5, -0.04, "Source: results.xlsx / Backbones; preliminary T4 latency, preprocessing excluded",
                ha="center", fontsize=9)
    (root / "curves").mkdir(exist_ok=True)
    figure.savefig(root / "curves/backbone_tradeoff.png", dpi=160, bbox_inches="tight")
    plt.close(figure)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("work_dir", type=Path)
    finalize_artifacts(parser.parse_args().work_dir)
